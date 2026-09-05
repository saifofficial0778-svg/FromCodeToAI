import os
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# ============================================================
# 1. GROQ SETUP
# ============================================================

load_dotenv()

api_key = os.getenv("GROQ_API_KEY")

if not api_key:
    raise ValueError("GROQ_API_KEY nahi mili")

client = Groq(api_key=api_key)

model = "openai/gpt-oss-120b"


# ============================================================
# 2. COMMON LLM FUNCTION
# ============================================================

def ask_llm(system_prompt, user_prompt):

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages
    )

    return response.choices[0].message.content


# ============================================================
# 3. JD DATA MODELS
# ============================================================

class ExperienceRequirement(BaseModel):
    min: int
    max: int


class JDData(BaseModel):
    required_skills: list[str]
    experience_years: ExperienceRequirement
    education: list[str] = Field(default_factory=list)


# ============================================================
# 4. STEP 1 — JD → JSON
# ============================================================

def step1_jd_extract(jd):

    print("\nSTEP 1: Extracting JD")

    system_prompt = """
    You are an HR assistant.

    Extract important requirements from the given Job Description.

    Return ONLY valid JSON.

    Required format:

    {
        "required_skills": [],
        "experience_years": {
            "min": 0,
            "max": 0
        },
        "education": []
    }

    Do not invent information.
    """

    user_prompt = f"""
    Extract the requirements from this Job Description:

    {jd}
    """

    return ask_llm(system_prompt, user_prompt)


# ============================================================
# 5. RESUME DATA MODELS
# ============================================================

class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Education(BaseModel):
    degree: str | None = None
    university: str | None = None
    year: str | None = None
class Project(BaseModel):
    name: str | None = None
    description: str | None = None
    technologies: list[str] = Field(default_factory=list)

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[Education] = Field(default_factory=list)
    projects: list[Project] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)

# ============================================================
# 6. PDF READER
# ============================================================

def read_pdf(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# ============================================================
# 7. DOCX READER
# ============================================================

def read_docx(file_path):

    document = Document(file_path)

    text = ""

    # Normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    # Tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():
                    text += cell.text + "\n"

    return text


# ============================================================
# 8. RESUME FILE READER
# ============================================================

def read_resume(file_path):

    file_path = Path(file_path)

    extension = file_path.suffix.lower()

    if extension == ".pdf":

        return read_pdf(file_path)

    elif extension == ".docx":

        return read_docx(file_path)

    else:

        raise ValueError(
            "Only PDF and DOCX files are supported"
        )


# ============================================================
# 9. STEP 2 — RESUME TEXT → JSON
# ============================================================

def step2_parse_resume(resume_text):

    print("\nSTEP 2: Parsing Resume")

    system_prompt = """
    You are an HR assistant.

    Extract important information from the candidate's resume.

    Return ONLY valid JSON.

    Required format:

    {
        "name": "",
        "email": "",
        "phone": "",
        "total_experience_years": 0,
        "skills": [],
        "experiences": [
            {
                "company": "",
                "role": "",
                "duration": "",
                "description": "",
                "skills_used": []
            }
        ],
        "education": [],
        "projects": [],
        "certifications": []
    }

    Do not invent information.
    """

    user_prompt = f"""
    Extract the information from this resume:

    {resume_text}
    """

    return ask_llm(system_prompt, user_prompt)


# ============================================================
# 10. MATCH RESULT MODEL
# ============================================================

class MatchResult(BaseModel):
    score: float
    matched_skills: list[str]
    missing_skills: list[str]
    experience_match: bool


# ============================================================
# 11. SKILL NORMALIZATION
# ============================================================
def normalize_skill(skill):

    skill = skill.lower().strip()

    replacements = {
        "react.js": "reactjs",
        "react js": "reactjs",

        "nodejs": "node.js",
        "node js": "node.js",

        "rest api": "restful api",
        "rest apis": "restful api",
        "restful api design": "restful api",
        "api development and consumption (restful apis)": "restful api",

        "version control (git)": "git",

        "authentication/authorization flows": "authentication",
        "jwt": "authentication",
        "rbac": "authorization",
    }

    return replacements.get(skill, skill)


# ============================================================
# 12. STEP 3 — COMPARE RESUME WITH JD
# ============================================================

def compare_resume_with_jd(jd_data, resume_data):

    print("\nSTEP 3: Comparing Resume with JD")

    # Normalize JD skills
    jd_skills = {
        normalize_skill(skill)
        for skill in jd_data.required_skills
    }

    resume_skills = {
        normalize_skill(skill)
        for skill in resume_data.skills
    }

    # Matched skills
    matched = jd_skills & resume_skills

    # Missing skills
    missing = jd_skills - resume_skills

    # Skill score
    if jd_skills:

        skill_score = (
            len(matched) / len(jd_skills)
        ) * 100

    else:

        skill_score = 0

    # Experience check
    if resume_data.total_experience_years is not None:

        experience_match = (
            resume_data.total_experience_years
            >= jd_data.experience_years.min
        )

    else:

        experience_match = False

    # Final score
    final_score = skill_score

    if experience_match:
        final_score += 10

    final_score = min(final_score, 100)

    return MatchResult(
        score=round(final_score, 2),
        matched_skills=sorted(matched),
        missing_skills=sorted(missing),
        experience_match=experience_match
    )


# ============================================================
# 13. JOB DESCRIPTION
# ============================================================

JD = """
About the job

Minimum Qualifications:

3–5 years of hands-on experience in fullstack development
with strong proficiency in ReactJS and Node.js.

Strong understanding of JavaScript/TypeScript fundamentals,
asynchronous programming, and modern frontend patterns.

Experience building and consuming APIs, handling state management,
and integrating frontend-backend workflows.

Working knowledge of testing practices, debugging,
and version control (Git).

Education: BE, ME, BTECH, MTECH, MCA, MSC, MBA
(or equivalent practical experience).

Key Responsibilities:

Design, develop, and ship end-to-end features using ReactJS
on the frontend and Node.js on the backend.

Develop RESTful APIs and backend modules in Node.js.

Implement authentication/authorization flows.

Write unit/integration tests.

Optimize frontend rendering and backend response times.

Contribute to AI-assisted features such as intelligent
recommendations, summarization, or workflow automation.
"""


# ============================================================
# 14. MAIN PIPELINE
# ============================================================

def main():

    # --------------------------------------------------------
    # STEP 1: JD → JSON → Pydantic
    # --------------------------------------------------------

    jd_json = step1_jd_extract(JD)

    jd_data = JDData.model_validate_json(jd_json)

    print("\nJD DATA:")
    print(jd_data)


    # --------------------------------------------------------
    # STEP 2: Resume File → Text
    # --------------------------------------------------------

    resume_path = Path("resumes/resume.pdf")

    resume_text = read_resume(resume_path)

    if not resume_text.strip():

        raise ValueError(
            "Resume se text extract nahi hua"
        )

    print("\nRESUME TEXT EXTRACTED")


    # --------------------------------------------------------
    # STEP 3: Resume Text → JSON → Pydantic
    # --------------------------------------------------------

    resume_json = step2_parse_resume(resume_text)

    resume_data = Resume.model_validate_json(resume_json)

    print("\nRESUME DATA:")
    print(resume_data)


    # --------------------------------------------------------
    # STEP 4: Compare
    # --------------------------------------------------------

    match_result = compare_resume_with_jd(
        jd_data,
        resume_data
    )


    # --------------------------------------------------------
    # STEP 5: Final Result
    # --------------------------------------------------------

    print("\n==============================")
    print("       FINAL HR RESULT")
    print("==============================")

    print(f"Score: {match_result.score}%")

    print("\nMatched Skills:")

    for skill in match_result.matched_skills:
        print(f"✓ {skill}")

    print("\nMissing Skills:")

    for skill in match_result.missing_skills:
        print(f"✗ {skill}")

    print(
        f"\nExperience Match: "
        f"{match_result.experience_match}"
    )


# ============================================================
# 15. RUN APPLICATION
# ============================================================

if __name__ == "__main__":
    main()